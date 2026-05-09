from config import PDF_SUPPORT, DOCX_SUPPORT
from text_utils import clean_extracted_text
from arabic_fixes import fix_arabic_pdf_text, fix_split_arabic_words, llm_correct_arabic_text

if PDF_SUPPORT:
    import pdfplumber
if DOCX_SUPPORT:
    from docx import Document
from pptx import Presentation

def extract_text_from_pdf(file):
    """Extract text from PDF file"""
    text = ""
    with pdfplumber.open(file) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    # Step 1: Fix reversed Arabic text from PDF visual-order extraction
    text = fix_arabic_pdf_text(text)
    # Step 1.5: Merge split Arabic words (e.g. 'الخدما ت' → 'الخدمات')
    text = fix_split_arabic_words(text)
    # Step 2: Use LLM to fix remaining extraction artifacts
    text = llm_correct_arabic_text(text)
    return clean_extracted_text(text)

def extract_text_from_docx(file):
    """Extract text from DOCX file"""
    doc = Document(file)
    text = ""
    for paragraph in doc.paragraphs:
        text += paragraph.text + "\n"
    return clean_extracted_text(text)

def extract_text_from_pptx(file):
    """Extract text from PPTX file"""
    prs = Presentation(file)
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return clean_extracted_text(text)

def extract_text(uploaded_file):
    """Extract text based on file type"""
    file_type = uploaded_file.name.split('.')[-1].lower()
    
    if file_type == 'pdf':
        if not PDF_SUPPORT:
            return None, "PDF support not installed. Run: pip install pdfplumber"
        return extract_text_from_pdf(uploaded_file), None
    elif file_type == 'docx':
        if not DOCX_SUPPORT:
            return None, "DOCX support not installed. Run: pip install python-docx"
        return extract_text_from_docx(uploaded_file), None
    elif file_type in ['pptx', 'ppt']:
        return extract_text_from_pptx(uploaded_file), None
    else:
        return None, f"Unsupported file type: {file_type}"
