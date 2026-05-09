import io
from datetime import datetime
from config import TRANSLATION_SUPPORT, LANG_CODES, RTL_LANGUAGES, DOCX_SUPPORT
from text_utils import clean_extracted_text

if TRANSLATION_SUPPORT:
    from deep_translator import GoogleTranslator
if DOCX_SUPPORT:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches

def translate_text(text, target_lang):
    """Translate text to target language"""
    if not TRANSLATION_SUPPORT:
        return "Translation not available. Install: pip install deep-translator"
    
    target_code = LANG_CODES.get(target_lang, "en")
    
    # Clean text before translation
    text = clean_extracted_text(text)
    
    # Split text into chunks (Google Translate has a 5000 char limit)
    max_chars = 4500
    chunks = [text[i:i+max_chars] for i in range(0, len(text), max_chars)]
    
    translated_chunks = []
    for chunk in chunks:
        if chunk.strip():
            try:
                translated = GoogleTranslator(source='auto', target=target_code).translate(chunk)
                translated_chunks.append(translated if translated else chunk)
            except Exception as e:
                translated_chunks.append(f"[Translation error: {str(e)}]")
    
    return "\n".join(translated_chunks)

def translate_single_text(text, target_code):
    """Translate a single text chunk"""
    if not text or not text.strip():
        return text
    if not TRANSLATION_SUPPORT:
        return text
    try:
        clean_text = text.strip()
        if len(clean_text) > 4500:
            clean_text = clean_text[:4500]
        translated = GoogleTranslator(source='auto', target=target_code).translate(clean_text)
        return translated if translated else text
    except Exception:
        return text

def translate_docx_inplace(file, target_lang):
    """Translate a DOCX file preserving original structure and formatting"""
    if not DOCX_SUPPORT:
        return None, "DOCX support not installed"
    
    target_code = LANG_CODES.get(target_lang, "en")
    
    # Load the document
    doc = Document(file)
    
    # Translate paragraphs while preserving formatting
    for para in doc.paragraphs:
        if para.text.strip():
            # Translate full paragraph text
            translated = translate_single_text(para.text, target_code)
            
            # Clear and set new text (preserves paragraph-level formatting)
            if para.runs:
                # Keep first run's formatting, update text
                first_run = para.runs[0]
                original_font = first_run.font.name
                original_size = first_run.font.size
                original_bold = first_run.font.bold
                original_italic = first_run.font.italic
                
                # Clear all runs
                for run in para.runs:
                    run.text = ""
                
                # Set translated text in first run
                first_run.text = translated
                first_run.font.name = original_font
                first_run.font.size = original_size
                first_run.font.bold = original_bold
                first_run.font.italic = original_italic
            else:
                para.text = translated
            
            # Set RTL alignment for RTL languages
            if target_lang in RTL_LANGUAGES:
                para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    # Translate tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if para.text.strip():
                        translated = translate_single_text(para.text, target_code)
                        if para.runs:
                            para.runs[0].text = translated
                            for run in para.runs[1:]:
                                run.text = ""
                        else:
                            para.text = translated
    
    # Save to buffer
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer, None

def translate_pptx_inplace(file, target_lang):
    """Translate a PPTX file preserving original structure and formatting"""
    target_code = LANG_CODES.get(target_lang, "en")
    
    # Load presentation
    prs = Presentation(file)
    
    # Translate each slide
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text_frame"):
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        if run.text.strip():
                            # Preserve formatting
                            original_font = run.font.name
                            original_size = run.font.size
                            original_bold = run.font.bold
                            original_italic = run.font.italic
                            
                            # Safely get color - handle scheme colors vs RGB
                            original_color = None
                            try:
                                if run.font.color and run.font.color.type is not None:
                                    try:
                                        original_color = run.font.color.rgb
                                    except AttributeError:
                                        # It's a scheme/theme color, skip RGB
                                        original_color = None
                            except Exception:
                                original_color = None
                            
                            # Translate
                            run.text = translate_single_text(run.text, target_code)
                            
                            # Restore formatting
                            run.font.name = original_font
                            run.font.size = original_size
                            run.font.bold = original_bold
                            run.font.italic = original_italic
                            if original_color:
                                run.font.color.rgb = original_color
            
            # Handle tables in slides
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text_frame:
                            for para in cell.text_frame.paragraphs:
                                for run in para.runs:
                                    if run.text.strip():
                                        run.text = translate_single_text(run.text, target_code)
    
    # Save to buffer
    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer, None

def create_translated_docx(translated_text, target_lang, original_filename):
    """Create a DOCX document from translated text"""
    doc = Document()
    
    # Add title
    doc.add_heading(f'Translated Document ({target_lang})', 0)
    
    # Add metadata
    doc.add_paragraph(f"Original file: {original_filename}")
    doc.add_paragraph(f"Translation date: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_paragraph("---")
    
    # Add translated content
    paragraphs = translated_text.split('\n')
    for para in paragraphs:
        if para.strip():
            p = doc.add_paragraph(para.strip())
            # Set RTL for RTL languages
            if target_lang in RTL_LANGUAGES:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    # Save to bytes
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def create_translated_pptx(translated_text, target_lang, original_filename):
    """Create a PPTX presentation from translated text"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Title slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = f"Translated Document ({target_lang})"
    subtitle.text = f"Original: {original_filename}\nDate: {datetime.now().strftime('%Y-%m-%d')}"
    
    # Content slides
    paragraphs = [p.strip() for p in translated_text.split('\n') if p.strip()]
    content_layout = prs.slide_layouts[1]
    
    # Group paragraphs into slides (max 6 paragraphs per slide)
    for i in range(0, len(paragraphs), 6):
        slide = prs.slides.add_slide(content_layout)
        title = slide.shapes.title
        title.text = f"Content (Page {i//6 + 1})"
        
        body = slide.placeholders[1]
        tf = body.text_frame
        
        for j, para in enumerate(paragraphs[i:i+6]):
            if j == 0:
                tf.text = para[:200]  # Limit text length
            else:
                p = tf.add_paragraph()
                p.text = para[:200]
    
    # Save to bytes
    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer
